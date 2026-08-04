import tkinter as tk
from tkinter import messagebox
from tkinter import ttk
import openpyxl
from datetime import datetime
import config

def mostrar_tela_emissao_nfe(janela_principal, limpar_tela, centralizar_janela, voltar_menu_callback):
    """Tela de Emissão e Gestão de NF-e preparada para Integração com API"""
    limpar_tela(janela_principal)
    centralizar_janela(janela_principal, 1100, 700)

    frame_principal = tk.Frame(janela_principal, padx=20, pady=15)
    frame_principal.pack(fill="both", expand=True)

    tk.Label(frame_principal, text="📄 Emissão de Nota Fiscal Eletrônica (NF-e)", font=("Arial", 15, "bold")).pack(anchor="w", pady=(0, 10))

    # --- PAINEL DE DADOS DA NOTA / DESTINATÁRIO ---
    f_form = tk.LabelFrame(frame_principal, text=" Dados da Nota & Destinatário ", font=("Arial", 9, "bold"), padx=15, pady=10)
    f_form.pack(fill="x", pady=(0, 10))

    f_l1 = tk.Frame(f_form)
    f_l1.pack(fill="x", pady=(0, 6))

    tk.Label(f_l1, text="Cód. Cliente:", font=("Arial", 9, "bold")).pack(side="left", padx=(0, 2))
    entry_cod_cli = tk.Entry(f_l1, font=("Arial", 10), width=12)
    entry_cod_cli.pack(side="left", padx=(0, 15))

    tk.Label(f_l1, text="Natureza da Operação:", font=("Arial", 9, "bold")).pack(side="left", padx=(0, 2))
    combo_nat_op = ttk.Combobox(f_l1, values=["Venda de Mercadoria", "Prestação de Serviço", "Remessa / Retorno"], width=22, state="readonly")
    combo_nat_op.set("Prestação de Serviço")
    combo_nat_op.pack(side="left", padx=(0, 15))

    tk.Label(f_l1, text="Ambiente:", font=("Arial", 9, "bold")).pack(side="left", padx=(0, 2))
    combo_ambiente = ttk.Combobox(f_l1, values=["Homologação (Testes)", "Produção (Valendo)"], width=20, state="readonly")
    combo_ambiente.set("Homologação (Testes)")
    combo_ambiente.pack(side="left")

    # --- PAINEL DE VALORES E DESCRIÇÃO ---
    f_l2 = tk.Frame(f_form)
    f_l2.pack(fill="x", pady=(4, 0))

    tk.Label(f_l2, text="Descrição / Serviços da NF-e:", font=("Arial", 9, "bold")).pack(side="left", padx=(0, 2))
    entry_desc_nfe = tk.Entry(f_l2, font=("Arial", 10), width=45)
    entry_desc_nfe.pack(side="left", padx=(0, 15))

    tk.Label(f_l2, text="Valor Total (R$):", font=("Arial", 9, "bold")).pack(side="left", padx=(0, 2))
    entry_valor_nfe = tk.Entry(f_l2, font=("Arial", 10), width=15)
    entry_valor_nfe.pack(side="left")
    entry_valor_nfe.insert(0, "0,00")

    # --- ÁREA DE LOGS / RETORNO DA API ---
    f_log_frame = tk.LabelFrame(frame_principal, text=" Retorno da API / Status de Transmissão ", font=("Arial", 9, "bold"), padx=10, pady=8)
    f_log_frame.pack(fill="both", expand=True, pady=(0, 10))

    txt_retorno_api = tk.Text(f_log_frame, font=("Courier New", 9), height=10, bg="#1E1E1E", fg="#00FFCC")
    txt_retorno_api.pack(side="left", fill="both", expand=True)
    txt_retorno_api.insert(tk.END, "[sistema] Módulo NF-e inicializado. Pronto para conexão com a API...\n")

    scrollbar_log = ttk.Scrollbar(f_log_frame, orient="vertical", command=txt_retorno_api.yview)
    txt_retorno_api.configure(yscrollcommand=scrollbar_log.set)
    scrollbar_log.pack(side="right", fill="y")

    # --- FUNÇÃO PRONTA PARA CONEXÃO COM API ---
    def transmitir_nfe_api():
        cod_cli = entry_cod_cli.get().strip()
        nat_op = combo_nat_op.get()
        ambiente = combo_ambiente.get()
        descricao = entry_desc_nfe.get().strip()
        valor = entry_valor_nfe.get().strip()

        if not cod_cli or not descricao:
            messagebox.showwarning("Atenção", "Preencha o Código do Cliente e a Descrição da Nota!")
            return

        txt_retorno_api.insert(tk.END, f"\n[{datetime.now().strftime('%H:%M:%S')}] Iniciando transmissão para a API ({ambiente})...\n")
        txt_retorno_api.see(tk.END)
        janela_principal.update()

        try:
            # =========================================================================
            # PONTO DE INTEGRAÇÃO COM A API (Substitua este bloco pela sua requisição HTTP)
            # Exemplo usando a biblioteca 'requests':
            # 
            # import requests
            # payload = {
            #     "cliente_cod": cod_cli,
            #     "natureza": nat_op,
            #     "descricao": descricao,
            #     "valor": valor,
            #     "ambiente": "homologacao" if "Homologação" in ambiente else "producao"
            # }
            # headers = {"Authorization": "Bearer SEU_TOKEN_DE_API"}
            # resposta = requests.post("https://api.suaempresa.com/v1/nfe", json=payload, headers=headers)
            # resultado = resposta.json()
            # =========================================================================

            # Simulação de resposta bem-sucedida da API para testes:
            import time
            time.sleep(1) # Simula o tempo de resposta da operadora/SEFAZ
            
            sucesso_simulado = True 

            if sucesso_simulado:
                protocolo_fake = f"PROT-{datetime.now().strftime('%Y%m%d%H%M%S')}"
                chave_fake = "35260700000000000000550010000000011234567890"
                
                txt_retorno_api.insert(tk.END, f"✅ [SUCESSO] Nota autorizada pela SEFAZ!\n")
                txt_retorno_api.insert(tk.END, f"   - Chave de Acesso: {chave_fake}\n")
                txt_retorno_api.insert(tk.END, f"   - Protocolo: {protocolo_fake}\n")
                txt_retorno_api.see(tk.END)
                
                messagebox.showinfo("Sucesso", f"Nota Fiscal emitida com sucesso!\nProtocolo: {protocolo_fake}")
                
                # Registra na auditoria do sistema (Aba 'Atualizacoes')
                pl = openpyxl.load_workbook(config.NOME_ARQUIVO)
                if "Atualizacoes" not in pl.sheetnames:
                    aba_at = pl.create_sheet("Atualizacoes")
                    aba_at.append(["Data", "Nº Conta", "Descrição do Log", "Item", "Quantidade", "Usuário"])
                else:
                    aba_at = pl['Atualizacoes']

                data_hoje = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
                aba_at.append([data_hoje, f"NF-e", f"Emissão de NF-e ({nat_op}) - R$ {valor}", descricao, 1, config.usuario_logado])
                pl.save(config.NOME_ARQUIVO)
                pl.close()

            else:
                txt_retorno_api.insert(tk.END, f"❌ [ERRO] Rejeição da SEFAZ: Erro de validação de schema.\n")
                txt_retorno_api.see(tk.END)
                messagebox.showerror("Erro de Emissão", "A API retornou uma rejeição. Verifique o log.")

        except Exception as e:
            txt_retorno_api.insert(tk.END, f"⚠️ [EXCEÇÃO] Falha de comunicação com a API: {str(e)}\n")
            txt_retorno_api.see(tk.END)
            messagebox.showerror("Erro de Conexão", f"Não foi possível conectar à API:\n{e}")

    # --- BOTÕES DE AÇÃO ---
    f_botoes = tk.Frame(frame_principal)
    f_botoes.pack(fill="x", pady=(5, 5))
    
    tk.Button(f_botoes, text="🚀 Transmitir NF-e para API", command=transmitir_nfe_api, bg="#38761D", fg="white", font=("Arial", 10, "bold"), height=2).pack(side="left", fill="x", expand=True, padx=(0, 5))
    tk.Button(frame_principal, text="⬅ Voltar ao Menu", command=voltar_menu_callback, bg="#595959", fg="white", font=("Arial", 10), height=1).pack(fill="x", pady=(5, 0))